from hashlib import sha256
from io import BytesIO
from pathlib import Path
from docx import Document as WordDocument
from openpyxl import load_workbook
from app.models.document import Document, Page
from .pdf_parser import extract_pdf_text, is_low_density
from .image_parser import parse_image
from .region import assess, doc_type_suffix
from app.parsers.ocr import OcrEngine, get_ocr_engine

MAX_BYTES = 8 * 1024 * 1024


def parse_document(file_name: str, content: bytes, ocr_engine: OcrEngine | None = None,
                   field_sink: dict | None = None) -> Document:
    """PRD §6.2 文件处理策略分发:类型识别 → 首选路径 → 回退路径(OCR)。

    任务④:OCR 路径同时评估困难区域(手写/低质),以 document_type 后缀标记
    (如 image_ocr_handwriting),供审核流程标记人工复核;不修改 Document 契约字段。

    ``field_sink``:发票专用 OCR 接口直接返回结构化字段时,引擎把它们放在
    ``engine.last_fields``;传入 sink 后由这里收集(键为项目字段名,值为规范化字符串),
    供上层优先采用(见 app/agent/nodes.py 的 extract 节点)。
    """
    if not content or len(content) > MAX_BYTES:
        raise ValueError("文件不能为空且不得超过8MB")
    engine = ocr_engine or get_ocr_engine()

    def collect_fields() -> None:
        if field_sink is None:
            return
        fields = getattr(engine, "last_fields", None) or {}
        if fields:
            field_sink.update(fields)
            setattr(engine, "last_fields", {})

    suffix = Path(file_name).suffix.lower()
    if suffix == ".pdf":
        # 首选:电子文本;回退:文本密度过低转 OCR(§6.2)
        pages = extract_pdf_text(content)
        if is_low_density(pages):
            text = engine.recognize(content)
            collect_fields()
            assessment = assess(text=text)   # PDF 字节无法做图片质量分析,仅评估文本质量
            pages = [Page(number=1, text=text)]
            doc_type = "pdf_ocr" + doc_type_suffix(assessment)
        else:
            doc_type = "pdf_text"
    elif suffix in {".png", ".jpg", ".jpeg"}:
        text = parse_image(content, engine)
        collect_fields()
        assessment = assess(content=content, text=text)
        pages = [Page(number=1, text=text)]
        doc_type = "image_ocr" + doc_type_suffix(assessment)
    elif suffix == ".txt":
        text = content.decode("utf-8-sig")
        assessment = assess(text=text)      # 文本材料同样检测手写/盖章关键词与乱码(§6.4)
        pages = [Page(number=1, text=text)]
        doc_type = "text" + doc_type_suffix(assessment)
    elif suffix == ".docx":
        word = WordDocument(BytesIO(content))
        text = [p.text for p in word.paragraphs]
        text += [" | ".join(c.text for c in row.cells) for t in word.tables for row in t.rows]
        joined = "\n".join(text)
        assessment = assess(text=joined)
        pages = [Page(number=1, text=joined)]
        doc_type = "docx" + doc_type_suffix(assessment)
    elif suffix == ".xlsx":
        book = load_workbook(BytesIO(content), read_only=True, data_only=True)
        text = [f"[{sheet.title}] " + " | ".join(
            " ".join(str(v) for v in row if v is not None)
            for row in sheet.iter_rows(values_only=True)) for sheet in book.worksheets]
        book.close()
        joined = "\n".join(text)
        assessment = assess(text=joined)
        pages = [Page(number=1, text=joined)]
        doc_type = "xlsx" + doc_type_suffix(assessment)
    else:
        raise ValueError("仅支持PDF、TXT、DOCX、XLSX及图片文件")
    return Document(document_id=sha256(content).hexdigest()[:16], file_name=Path(file_name).name,
                    document_type=doc_type, pages=pages)
